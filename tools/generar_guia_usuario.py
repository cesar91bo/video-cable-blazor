from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Documentacion" / "Guia_funcional_usuarios_Video_Cable.docx"
LOGO = ROOT / "VideoCable.Web" / "wwwroot" / "images" / "logo.jpg"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GRAY = "5F6B76"
WHITE = "FFFFFF"
INK = "24313D"

def set_font(run, name="Calibri", size=11, bold=False, color=INK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=100, start=140, bottom=100, end=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_cell_text(cell, text, bold=False, color=INK, size=10.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)

def set_table_geometry(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(width)); grid.append(gc)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(width)); tcW.set(qn("w:type"), "dxa")

def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Página "); set_font(r, size=9, color=GRAY)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    rr = paragraph.add_run(); rr._r.extend([begin, instr, end]); set_font(rr, size=9, color=GRAY)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p

def add_body(doc, text, bold_start=None):
    p = doc.add_paragraph(style="Normal")
    if bold_start and text.startswith(bold_start):
        set_font(p.add_run(bold_start), bold=True)
        set_font(p.add_run(text[len(bold_start):]))
    else:
        set_font(p.add_run(text))
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_font(p.add_run(text))
    return p

def add_step(doc, title, detail):
    p = doc.add_paragraph(style="List Number")
    set_font(p.add_run(title + ". "), bold=True, color=NAVY)
    set_font(p.add_run(detail))
    return p

def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    cell = table.cell(0,0); shade(cell, PALE); margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label + ": "), bold=True, color=BLUE)
    set_font(p.add_run(text))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in (("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,NAVY,10,5)):
    s=styles[name]; s.font.name="Calibri"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
for name in ("List Bullet", "List Bullet 2", "List Number"):
    s=styles[name]; s.font.name="Calibri"; s.font.size=Pt(11); s.paragraph_format.space_after=Pt(4); s.paragraph_format.line_spacing=1.25
styles["List Bullet"].paragraph_format.left_indent=Inches(.375); styles["List Bullet"].paragraph_format.first_line_indent=Inches(-.188)
styles["List Bullet 2"].paragraph_format.left_indent=Inches(.65); styles["List Bullet 2"].paragraph_format.first_line_indent=Inches(-.188)
styles["List Number"].paragraph_format.left_indent=Inches(.375); styles["List Number"].paragraph_format.first_line_indent=Inches(-.188)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("VIDEO CABLE  |  GUÍA FUNCIONAL"), size=8.5, bold=True, color=GRAY)
add_page_field(sec.footer.paragraphs[0])

# Cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(34); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
if LOGO.exists(): p.add_run().add_picture(str(LOGO), width=Inches(1.35))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(40); p.paragraph_format.space_after=Pt(10)
set_font(p.add_run("GUÍA FUNCIONAL"), size=12, bold=True, color=BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10)
set_font(p.add_run("Sistema de Gestión Video Cable"), size=28, bold=True, color=NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(34)
set_font(p.add_run("Descripción general para usuarios"), size=15, color=GRAY)
add_callout(doc, "Propósito", "Explicar de manera simple qué permite hacer el sistema y cómo se relacionan sus principales módulos de trabajo.")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(120)
set_font(p.add_run("Documento elaborado a partir de las funciones disponibles en la aplicación"), size=9.5, italic=True, color=GRAY)
doc.add_page_break()

add_heading(doc, "1. ¿Qué es el sistema?", 1)
add_body(doc, "Video Cable es una aplicación de gestión administrativa para organizar la operación diaria de una empresa que brinda servicios de televisión, internet u otros planes asociados. Centraliza la información de clientes, suscripciones, facturación, cobranzas, caja, reclamos y deudas en un mismo lugar.")
add_body(doc, "El eje del sistema es la suscripción: vincula a un cliente con un servicio, una ubicación de instalación y, cuando corresponde, una caja de distribución. A partir de ella se generan las facturas periódicas, se consulta la deuda y se registra el cobro.")
add_callout(doc, "En pocas palabras", "El sistema acompaña el circuito completo: configurar servicios, dar de alta clientes, asociar suscripciones, facturar, cobrar, controlar la caja y consultar pendientes.")

add_heading(doc, "2. Recorrido general de trabajo", 1)
for t,d in [
    ("Preparar la configuración", "cargar los datos de la empresa, los servicios o planes y las cajas de distribución"),
    ("Registrar al cliente", "guardar sus datos de identificación, contacto, domicilio y condición impositiva"),
    ("Crear la suscripción", "asignar un servicio, precio, estado, caja y dirección de instalación"),
    ("Generar la facturación", "crear las facturas del período para las suscripciones que corresponda"),
    ("Consultar y cobrar", "ver el estado de cuenta, registrar pagos totales o parciales y emitir un comprobante"),
    ("Controlar la operación", "abrir y cerrar caja, gestionar reclamos y revisar reportes de deuda")]: add_step(doc,t,d)

add_heading(doc, "3. Pantalla de inicio", 1)
add_body(doc, "El Panel de control resume los datos más importantes de la jornada y del período actual. Permite ver rápidamente el estado de la caja diaria, lo cobrado hoy, la cantidad de facturas pendientes y el total facturado en el mes.")
add_bullet(doc, "Ingresos, egresos y saldo de la caja del día.")
add_bullet(doc, "Importe cobrado durante la jornada.")
add_bullet(doc, "Facturas que continúan pendientes.")
add_bullet(doc, "Facturación acumulada del mes actual.")
add_body(doc, "El botón de actualización vuelve a consultar la información para mostrar los valores más recientes.")

doc.add_page_break()
add_heading(doc, "4. Funciones de administración", 1)

add_heading(doc, "4.1 Configuración de la empresa", 2)
add_body(doc, "Permite registrar y actualizar los datos generales y fiscales que identifican a la empresa: nombre de fantasía, razón social, CUIT, condición frente al IVA, Ingresos Brutos, domicilio, localidad, provincia, teléfono y correo electrónico.")
add_body(doc, "Esta información se utiliza como referencia institucional y puede aparecer en comprobantes o procesos de facturación.")

add_heading(doc, "4.2 Servicios y planes", 2)
add_body(doc, "En este módulo se mantiene el catálogo comercial. Cada servicio puede tener un código, una descripción, un tipo, una modalidad de cobro, un precio vigente y un estado activo o inactivo.")
add_bullet(doc, "Crear nuevos planes o conceptos de servicio.")
add_bullet(doc, "Modificar descripción, clasificación y precio actual.")
add_bullet(doc, "Definir si el cobro es mensual o corresponde a otro tipo de concepto.")
add_bullet(doc, "Desactivar servicios que ya no se ofrecen sin perder su referencia histórica.")

add_heading(doc, "4.3 Cajas de distribución", 2)
add_body(doc, "Permite identificar y ubicar las cajas o puntos de distribución de la red. Se registra código, descripción, dirección o referencia, zona, coordenadas geográficas y estado activo.")
add_bullet(doc, "Editar los datos de cada caja.")
add_bullet(doc, "Visualizar su ubicación en un mapa.")
add_bullet(doc, "Consultar un resumen de cajas y los clientes vinculados a cada una.")

add_heading(doc, "5. Gestión de clientes y suscripciones", 1)
add_heading(doc, "5.1 Clientes", 2)
add_body(doc, "La pantalla Clientes concentra el padrón de abonados. La búsqueda admite nombre, código, número de documento o CUIT. Para cada cliente pueden cargarse datos personales o comerciales, condición impositiva, domicilio, localidad, teléfono, correo y observaciones.")
add_bullet(doc, "Crear y editar clientes.")
add_bullet(doc, "Marcar un cliente como activo o inactivo.")
add_bullet(doc, "Evitar búsquedas manuales extensas mediante filtros rápidos.")

add_heading(doc, "5.2 Suscripciones", 2)
add_body(doc, "Una suscripción representa el servicio contratado por un cliente. Al crearla se seleccionan el cliente, el plan mensual, la caja de distribución y la fecha de alta. También se define el estado, el precio vigente y la dirección específica de instalación.")
add_body(doc, "La aplicación contempla estados Activo, Suspendido y Baja. También permite guardar fecha de baja, localidad, referencias para llegar al domicilio y observaciones operativas.")
add_callout(doc, "Importante", "Un mismo cliente puede ser localizado por sus datos y luego asociado a la suscripción correspondiente. La dirección de instalación puede diferir del domicilio principal del titular.")

doc.add_page_break()
add_heading(doc, "6. Facturación y estado de cuenta", 1)
add_heading(doc, "6.1 Facturación mensual", 2)
add_body(doc, "Este módulo genera las facturas internas correspondientes a un mes y año. El sistema toma las suscripciones aplicables y evita volver a crear una factura ya existente para el mismo período.")
add_bullet(doc, "Seleccionar mes y año a procesar.")
add_bullet(doc, "Generar la facturación del período.")
add_bullet(doc, "Consultar las facturas ya creadas.")
add_bullet(doc, "Ver el resultado del proceso: generadas, omitidas y total de suscripciones revisadas.")
add_bullet(doc, "Imprimir una factura o gestionar acciones disponibles sobre ella.")
add_body(doc, "También existe una carga manual para registrar una factura puntual o una deuda proveniente de un sistema anterior. En ese caso se seleccionan cliente y suscripción, se indica el concepto, importe, fechas y observaciones.")

add_heading(doc, "6.2 Estado de cuenta", 2)
add_body(doc, "Permite consultar la situación de un cliente en particular. Presenta sus facturas pendientes, la cantidad de meses adeudados y el total a pagar. Desde la misma pantalla se puede imprimir el detalle o continuar directamente hacia Cobranza.")

add_heading(doc, "6.3 Impresión de facturas", 2)
add_body(doc, "La vista de impresión muestra los datos de la empresa, el tipo de comprobante, los datos del cliente, los conceptos facturados y los importes. Cuando la información fiscal todavía no está autorizada, el sistema lo indica en pantalla para evitar confusiones.")

add_heading(doc, "7. Cobranza", 1)
add_body(doc, "Cobranza es la pantalla operativa para registrar pagos. Primero se busca y selecciona el cliente; luego se cargan sus facturas con saldo pendiente y el total adeudado.")
for t,d in [
    ("Buscar al cliente", "por nombre, código, documento o CUIT"),
    ("Revisar la deuda", "ver facturas pendientes y saldos disponibles para cobrar"),
    ("Indicar el monto", "cobrar el total o ingresar un pago parcial"),
    ("Completar el pago", "seleccionar medio de pago y agregar referencia u observaciones si corresponde"),
    ("Confirmar", "registrar el pago y, opcionalmente, imprimir el comprobante")]: add_step(doc,t,d)
add_body(doc, "Los medios previstos son efectivo, transferencia, tarjeta u otro. El pago se aplica a las facturas pendientes y el saldo de cada una queda actualizado como pendiente, parcial o pagado.")
add_callout(doc, "Requisito operativo", "Para registrar cobranzas debe existir una caja diaria abierta. Esta condición asegura que el ingreso quede reflejado en el control de caja.")

add_heading(doc, "7.1 Comprobante de pago", 2)
add_body(doc, "Después de cobrar, el sistema puede presentar un comprobante imprimible con los datos de la empresa, el pago realizado, el cliente y el detalle de la operación.")

doc.add_page_break()
add_heading(doc, "8. Caja diaria", 1)
add_body(doc, "La caja diaria permite controlar el dinero y los movimientos de cada jornada.")
add_heading(doc, "Apertura", 2)
add_body(doc, "Al comenzar el día se informa el monto inicial y, si hace falta, una observación. Solo puede existir una caja abierta a la vez.")
add_heading(doc, "Movimientos", 2)
add_body(doc, "Los cobros registrados desde Cobranza generan movimientos de caja. Además, el usuario puede cargar manualmente ingresos, egresos o ajustes, indicando importe y concepto.")
add_heading(doc, "Cierre", 2)
add_body(doc, "Al finalizar la jornada se informa el monto final contado y una observación de cierre. El sistema compara ese valor con el saldo calculado y conserva la diferencia para control.")

add_heading(doc, "9. Reclamos", 1)
add_body(doc, "El módulo Reclamos organiza la atención de incidencias de clientes. Los registros se pueden buscar por cliente o texto y filtrar por estado, tipo y prioridad.")
add_bullet(doc, "Relacionar el reclamo con un cliente, una suscripción y una caja de distribución.")
add_bullet(doc, "Registrar fecha, tipo, prioridad y descripción del problema.")
add_bullet(doc, "Agregar observaciones internas para operadores o técnicos.")
add_bullet(doc, "Documentar la resolución aplicada.")
add_bullet(doc, "Cambiar el estado a En proceso, Resuelto o Cerrado.")
add_body(doc, "Los indicadores superiores muestran cuántos reclamos están abiertos, en proceso, resueltos y cuántos integran el listado actual.")

add_heading(doc, "10. Reporte de deudas", 1)
add_body(doc, "Ofrece una visión general de los clientes con deuda. Puede filtrarse por cliente, ordenarse según diferentes criterios y limitarse por antigüedad.")
add_bullet(doc, "Cantidad de clientes con deuda.")
add_bullet(doc, "Monto total adeudado.")
add_bullet(doc, "Cantidad de facturas pendientes.")
add_bullet(doc, "Promedio de deuda por cliente.")
add_body(doc, "Desde el reporte se puede imprimir el resultado actual o pasar a Cobranza para atender un caso seleccionado.")

add_heading(doc, "11. Presupuestos", 1)
add_body(doc, "Permite confeccionar e imprimir presupuestos sin generar una factura. Se puede seleccionar un cliente existente o completar un destinatario, fecha, dirección, teléfono, vigencia, número y observaciones.")
add_body(doc, "El usuario agrega uno o más ítems con descripción, cantidad y precio unitario. El sistema calcula los importes y prepara una presentación imprimible.")

doc.add_page_break()
add_heading(doc, "12. Relación entre los módulos", 1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
headers=["Información", "Se origina en", "Se utiliza en"]
for i,h in enumerate(headers): set_cell_text(table.rows[0].cells[i], h, True, WHITE); shade(table.rows[0].cells[i], NAVY)
for row in [
    ("Cliente", "Clientes", "Suscripciones, facturación, cobranza, reclamos y presupuestos"),
    ("Servicio / plan", "Servicios / Planes", "Suscripciones y conceptos facturados"),
    ("Caja de distribución", "Cajas de distribución", "Suscripciones, mapa y reclamos"),
    ("Suscripción", "Suscripciones", "Facturación mensual, deuda y cobranza"),
    ("Factura", "Facturación mensual", "Estado de cuenta, reporte de deudas y cobranza"),
    ("Pago", "Cobranza", "Comprobante, saldo de facturas y caja diaria"),
    ("Movimiento", "Cobranza o Caja diaria", "Control y cierre de la jornada")]:
    cells=table.add_row().cells
    for i,v in enumerate(row): set_cell_text(cells[i],v,bold=(i==0)); shade(cells[i], WHITE if len(table.rows)%2 else LIGHT)
set_repeat_header(table.rows[0]); set_table_geometry(table,[2100,2400,4860])

add_heading(doc, "13. Recomendaciones de uso", 1)
add_bullet(doc, "Mantener actualizados los datos de empresa, clientes, servicios y cajas.")
add_bullet(doc, "Abrir la caja antes de comenzar a cobrar y cerrarla al terminar la jornada.")
add_bullet(doc, "Revisar el período antes de generar facturación mensual.")
add_bullet(doc, "Usar observaciones y referencias para dejar trazabilidad de casos especiales.")
add_bullet(doc, "No eliminar información histórica: utilizar estados inactivo, suspendido o baja cuando corresponda.")
add_bullet(doc, "Consultar el estado de cuenta antes de informar una deuda y verificar el comprobante luego del cobro.")

add_heading(doc, "14. Alcance y consideraciones", 1)
add_body(doc, "La guía describe las funciones visibles actualmente en el menú principal y su comportamiento operativo. Algunas capacidades fiscales asociadas con ARCA, certificados, punto de venta, CAE y comprobantes electrónicos forman parte de la configuración técnica o de una etapa de integración. Por ese motivo, no se presentan aquí como tareas habituales del usuario general.")
add_body(doc, "Las pantallas de prueba técnica y las páginas internas de diagnóstico tampoco forman parte del circuito normal de uso.")
add_callout(doc, "Resumen final", "Video Cable reúne en una sola aplicación la administración de abonados y servicios, la facturación periódica, el seguimiento de deuda, la cobranza, el control diario de caja y la gestión de reclamos.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
