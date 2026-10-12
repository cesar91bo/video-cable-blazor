from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

root = Path(__file__).resolve().parents[1]
source = root / 'docs/manual_usuario.md'
output = root / 'docs/Manual de usuario Video Cable.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.8), Inches(.8)
sec.header_distance, sec.footer_distance = Inches(.28), Inches(.28)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(11)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.05
for name, size, color in [('Title',26,'000000'),('Heading 1',20,'17364D'),('Heading 2',13,'17364D')]:
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.color.rgb = 'Calibri', Pt(size), RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(10 if name=='Heading 2' else 0)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['List Bullet','List Number']:
    doc.styles[name].font.size = Pt(11)
    doc.styles[name].paragraph_format.space_after = Pt(5)
    doc.styles[name].paragraph_format.line_spacing = 1.05
header = sec.header.paragraphs[0]
header.text = 'VIDEO CABLE  |  Manual de usuario'
header.style = 'Normal'
header.runs[0].font.size = Pt(9)
header.runs[0].font.color.rgb = RGBColor.from_string('536571')
footer = sec.footer.paragraphs[0]
footer.text = 'Versión 1.0  •  12 octubre 2026' + '\t' + 'Página '
footer.paragraph_format.tab_stops.add_tab_stop(Inches(5.9))
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(9)
pages = re.split(r'(?m)^# ',source.read_text(encoding='utf-8'))[1:]
for index, page in enumerate(pages):
    lines = page.strip().splitlines()
    chapter = doc.add_paragraph(lines[0], 'Title' if index == 0 else 'Heading 1')
    if index: chapter.paragraph_format.page_break_before = True
    for line in lines[1:]:
        if not line.strip(): continue
        if line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 2')
        elif re.match(r'^\d+\. ',line):
            # Explicit numbering restarts for each procedure and survives PDF rendering.
            p = doc.add_paragraph(line)
            p.paragraph_format.left_indent = Inches(.18)
            p.paragraph_format.first_line_indent = Inches(-.18)
            p.paragraph_format.space_after = Pt(5)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:],'List Bullet')
        else:
            p = doc.add_paragraph(line)
            if index == 0 and re.match(r'^\d+ — ',line):
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(10)
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
doc.core_properties.title = 'Manual de usuario de Video Cable'
doc.core_properties.subject = 'Operación del sistema y registro de capacitación'
doc.core_properties.author = 'Video Cable'
doc.core_properties.keywords = 'manual usuario capacitación facturación cobranza'
doc.save(output)
print(output)
print(f'Capítulos previstos: {len(pages)}')
